"""Create a small canonical BrEaST snapshot directly from original workbook and PNG files."""

import argparse
import hashlib
import json
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from medbenchagent.data.dataset import (
    DatasetManifest,
    FileDatasetReader,
    file_digest,
    seal_manifest,
)
from medbenchagent.data.models import AssetEntry, CaseRecord, Evidence, SourceRef


def prepare(workbook: Path, image_root: Path, output: Path, limit: int = 20):
    if limit < 1:
        raise ValueError("limit must be positive")
    if output.exists():
        raise ValueError("Use a new output directory; snapshots are immutable")
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        sheet = book["BrEaST-Lesions-USG clinical dat"]
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        records = [dict(zip(headers, row, strict=True)) for row in rows if row[0] is not None]
        dictionary = [
            {
                "field": row[0],
                "description": row[1],
                "source": {
                    "source_id": "metadata",
                    "path": "source.xlsx",
                    "record": f"Data Dictionary!{index}",
                },
            }
            for index, row in enumerate(
                book["Data Dictionary"].iter_rows(values_only=True), start=1
            )
            if index > 1 and row[0] is not None
        ]
    finally:
        book.close()
    records.sort(key=lambda row: int(row["CaseID"]))
    records = records[:limit]
    if not records or len({row["CaseID"] for row in records}) != len(records):
        raise ValueError("Expected unique, nonempty source CaseIDs")
    images = {}
    for path in image_root.rglob("*.png"):
        if path.name in images:
            raise ValueError(f"Ambiguous image filename: {path.name}")
        images[path.name] = path
    output.mkdir(parents=True)
    metadata = output / "sources" / "metadata"
    pixels = output / "sources" / "images"
    metadata.mkdir(parents=True)
    pixels.mkdir(parents=True)
    shutil.copyfile(workbook, metadata / "source.xlsx")
    assets, cases = [], []
    for row in records:
        case_id = f"case-{int(row['CaseID']):03d}"
        evidence = []

        def source(field):
            return SourceRef(
                source_id="metadata", path="source.xlsx", record=str(row["CaseID"]), field=field
            )

        for column, visible in (
            ("Image_filename", True),
            ("Mask_tumor_filename", False),
            ("Mask_other_filename", False),
        ):
            names = str(row[column]).split("&") if row[column] else []
            for index, name in enumerate(names):
                name = name.strip()
                path = images[name]
                shutil.copyfile(path, pixels / name)
                with Image.open(path) as image:
                    width, height = image.size
                asset_id = f"{case_id}-{column.lower()}-{index}"
                assets.append(
                    AssetEntry(
                        asset_id=asset_id,
                        source_id="images",
                        source_path=name,
                        file_digest=file_digest(path),
                        kind="image",
                        format="png",
                        width=width,
                        height=height,
                        metadata={"semantic_role": "primary" if visible else "mask"},
                    )
                )
                evidence.append(
                    Evidence(
                        id="image" if visible else f"{column.lower()}-{index}",
                        name="image" if visible else column.lower(),
                        type="asset",
                        value=asset_id,
                        ref=[] if visible else ["image"],
                        source_refs=[source(column)],
                        input_visibility="model_visible" if visible else "reference_only",
                    )
                )
        for column, value in row.items():
            if column in {"CaseID", "Image_filename", "Mask_tumor_filename", "Mask_other_filename"}:
                continue
            name = re.sub(r"[^a-z0-9]+", "_", column.lower()).strip("_")
            evidence.append(
                Evidence(
                    id=name,
                    name=name,
                    type="value",
                    value=value,
                    ref=["image"],
                    source_refs=[source(column)],
                    input_visibility="reference_only",
                )
            )
        cases.append(CaseRecord(case_id=case_id, patient_id=None, evidence=evidence))
    for name, rows in (("cases.jsonl", cases), ("assets.jsonl", assets)):
        (output / name).write_text(
            "".join(row.model_dump_json() + "\n" for row in rows), encoding="utf-8"
        )
    mapping = {
        "case_unit": "One source lesion record and its named ultrasound image; no patient/study grouping inferred",
        "case_id": "case- plus source CaseID padded to three digits",
        "patient_id": "unknown",
        "study_id": "unknown",
        "source_sheet": sheet.title,
        "field_mapping": {
            name.lower(): name
            for name in headers
            if name
            not in {"CaseID", "Image_filename", "Mask_tumor_filename", "Mask_other_filename"}
        },
        "data_dictionary": dictionary,
        "input_policy": "Only image is model_visible; annotations and masks are reference_only",
        "selection": {"order": "ascending numeric CaseID", "limit": limit, "actual": len(cases)},
        "derivations": "None; original values and masks preserved, no bounding boxes inferred",
    }
    (output / "mapping.json").write_text(
        json.dumps(mapping, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    manifest = seal_manifest(
        output,
        DatasetManifest(
            dataset_id="BrEaST",
            revision="pending",
            cases_path="cases.jsonl",
            assets_path="assets.jsonl",
            mapping_path="mapping.json",
            sources=[
                {"source_id": "metadata", "root": "sources/metadata"},
                {"source_id": "images", "root": "sources/images"},
            ],
        ),
    )
    manifest.revision = hashlib.sha256(
        json.dumps([item.model_dump() for item in manifest.members], sort_keys=True).encode()
    ).hexdigest()
    path = output / "manifest.json"
    path.write_text(manifest.model_dump_json(indent=2) + "\n", encoding="utf-8")
    FileDatasetReader(path)
    return path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--image-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--limit", type=int, default=20)
    args = parser.parse_args()
    print(
        prepare(
            args.workbook.resolve(), args.image_root.resolve(), args.output.resolve(), args.limit
        )
    )


if __name__ == "__main__":
    main()
