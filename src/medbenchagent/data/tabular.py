"""Streaming table inspection with explicit scan coverage and identity checks."""

import csv
from collections import Counter
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path


@contextmanager
def _rows(path, sheet):
    if path.suffix.lower() == ".csv":
        if sheet is not None:
            raise ValueError("CSV has no worksheets")
        with path.open(encoding="utf-8-sig", newline="") as handle:
            yield iter(csv.reader(handle)), None
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            selected = book.worksheets[0] if sheet is None else book[sheet]
            yield iter(selected.iter_rows(values_only=True)), selected.title
        finally:
            book.close()
    else:
        raise ValueError("Table inspection supports CSV and XLSX")


def inspect_table(path, *, sheet=None, id_columns=(), preview_rows=5, max_rows=None):
    """Measure scanned rows; max_rows never silently represents full-table coverage."""
    if (
        type(preview_rows) is not int
        or preview_rows < 0
        or (max_rows is not None and (type(max_rows) is not int or max_rows < 1))
    ):
        raise ValueError("Invalid table scan limits")
    path = Path(path)
    with _rows(path, sheet) as (rows, selected):
        headers = next(rows, None)
        if (
            not headers
            or any(not isinstance(h, str) or not h.strip() for h in headers)
            or len(set(headers)) != len(headers)
        ):
            raise ValueError("Table requires distinct nonempty column names")

        if len(id_columns) != len(set(id_columns)) or set(id_columns) - set(headers):
            raise ValueError("ID columns must select distinct table columns")
        missing, kinds = Counter(), {h: Counter() for h in headers}
        identities, duplicate_rows, missing_ids = set(), 0, 0
        preview, count, complete = [], 0, True
        for values in rows:
            if max_rows is not None and count == max_rows:
                complete = False
                break
            if len(values) != len(headers):
                raise ValueError(f"Row {count + 2} does not match column count")
            row = dict(zip(headers, values, strict=True))
            for name, value in row.items():
                if value is None or value == "":
                    missing[name] += 1
                else:
                    kinds[name][type(value).__name__] += 1
            if id_columns:
                identity = tuple(row[name] for name in id_columns)
                if any(value is None or value == "" for value in identity):
                    missing_ids += 1
                else:
                    duplicate_rows += identity in identities
                    identities.add(identity)
            if len(preview) < preview_rows:
                preview.append(
                    {
                        name: value.isoformat() if isinstance(value, (date, datetime)) else value
                        for name, value in row.items()
                    }
                )
            count += 1
    return {
        "path": str(path.resolve()),
        "sheet": selected,
        "rows_scanned": count,
        "complete": complete,
        "columns": [{"name": h, "missing": missing[h], "types": dict(kinds[h])} for h in headers],
        "id_columns": list(id_columns),
        "duplicate_id_rows": duplicate_rows,
        "missing_id_rows": missing_ids,
        "preview": preview,
    }
